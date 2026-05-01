"""Tests for ``PdfLoader`` and ``DocxLoader``.

Fixtures are built inline at test time rather than committed to the repo —
- ``sample_pdf``: hand-crafted minimal text PDF (~600 bytes)
- ``scanned_pdf``: ``pypdf.PdfWriter().add_blank_page()`` (no extractable text)
- ``sample_docx``: ``python-docx`` Document with H1+H2 headings
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from mdwiki.loaders import get_loader_for
from mdwiki.loaders.docx import DocxLoader
from mdwiki.loaders.pdf import PdfLoader


def _make_text_pdf(text: str) -> bytes:
    """Construct a minimal valid 1-page PDF containing literal text via Helvetica.

    Used as a fixture builder so tests don't need a committed binary file.
    """
    content = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(content)).encode() + b">>\nstream\n" + content + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    pdf = bytearray(header)
    offsets: list[int] = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{i} 0 obj\n".encode())
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")
    xref_off = len(pdf)
    pdf.extend(b"xref\n")
    pdf.extend(f"0 {len(objs) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for off in offsets:
        pdf.extend(f"{off:010d} 00000 n \n".encode())
    pdf.extend(b"trailer\n")
    pdf.extend(f"<</Size {len(objs) + 1}/Root 1 0 R>>\n".encode())
    pdf.extend(b"startxref\n")
    pdf.extend(f"{xref_off}\n".encode())
    pdf.extend(b"%%EOF\n")
    return bytes(pdf)


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """A 1-page PDF containing a known sentence pypdf can extract."""
    src = tmp_path / "sample.pdf"
    src.write_bytes(_make_text_pdf("expected sentence from sample.pdf"))
    return src


@pytest.fixture
def scanned_pdf(tmp_path: Path) -> Path:
    """A 1-page PDF with NO extractable text (image-only / scanned simulation)."""
    src = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    src.write_bytes(buf.getvalue())
    return src


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """A DOCX with H1, H2, paragraphs, and a small table."""
    src = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Top-Level Title", level=1)
    doc.add_paragraph("First paragraph under the H1.")
    doc.add_heading("Second-Level Section", level=2)
    doc.add_paragraph("Body paragraph under the H2.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "header-a"
    table.rows[0].cells[1].text = "header-b"
    table.rows[1].cells[0].text = "value-a"
    table.rows[1].cells[1].text = "value-b"
    doc.save(str(src))
    return src


# ----- PdfLoader -----


@pytest.mark.unit
def test_pdf_loader_extracts_text_from_text_pdf(sample_pdf: Path) -> None:
    """``PdfLoader`` returns markdown containing the PDF's literal text."""
    md = PdfLoader().load_to_markdown(sample_pdf)
    assert "expected sentence from sample.pdf" in md


@pytest.mark.unit
def test_pdf_loader_warns_on_empty_extraction(scanned_pdf: Path, caplog: pytest.LogCaptureFixture) -> None:
    """A scanned PDF (no extractable text) returns empty + logs a 'no text' warning."""
    with caplog.at_level(logging.WARNING, logger="mdwiki.loaders.pdf"):
        md = PdfLoader().load_to_markdown(scanned_pdf)
    assert md.strip() == ""
    assert any("no text" in r.message.lower() for r in caplog.records), [r.message for r in caplog.records]


@pytest.mark.unit
def test_pdf_loader_can_handle_pdf_extension(tmp_path: Path) -> None:
    src = tmp_path / "any.pdf"
    src.write_bytes(b"%PDF-1.4 fake")
    loader = PdfLoader()
    assert loader.can_handle(src) is True
    assert loader.can_handle(Path("not.txt")) is False


@pytest.mark.unit
def test_pdf_loader_routed_via_registry(sample_pdf: Path) -> None:
    """``get_loader_for(pdf)`` returns ``PdfLoader``."""
    assert isinstance(get_loader_for(sample_pdf), PdfLoader)


@pytest.mark.unit
def test_pdf_loader_vision_fallback_calls_describe_image_when_extraction_empty(scanned_pdf: Path, mocker: object) -> None:
    """vision_fallback=True + empty extraction → render pages, call provider.describe_image, return its output."""
    fake_provider = mocker.Mock()
    fake_provider.describe_image.return_value = "## Page 1\n\nA scanned page with diagram."

    loader = PdfLoader(provider=fake_provider, vision_fallback=True)
    md = loader.load_to_markdown(scanned_pdf)

    assert "scanned page" in md.lower()
    fake_provider.describe_image.assert_called()
    # The first arg is a path to a rendered PNG
    call_path = fake_provider.describe_image.call_args.args[0]
    assert isinstance(call_path, Path)
    assert call_path.suffix.lower() == ".png"


@pytest.mark.unit
def test_pdf_loader_vision_fallback_disabled_falls_back_to_warning(scanned_pdf: Path, caplog: pytest.LogCaptureFixture) -> None:
    """vision_fallback=False (default) keeps the v1.0 behavior: empty + warning, no provider call."""
    import logging

    fake_provider = type("X", (), {"describe_image": lambda self, p: (_ for _ in ()).throw(AssertionError("must not call"))})()
    with caplog.at_level(logging.WARNING, logger="mdwiki.loaders.pdf"):
        md = PdfLoader(provider=fake_provider, vision_fallback=False).load_to_markdown(scanned_pdf)
    assert md == ""
    assert any("no text" in r.message.lower() for r in caplog.records)


@pytest.mark.unit
def test_pdf_loader_vision_fallback_skipped_when_extraction_succeeds(sample_pdf: Path, mocker: object) -> None:
    """If pypdf extraction yields text, vision_fallback is not triggered (no provider call)."""
    fake_provider = mocker.Mock()
    md = PdfLoader(provider=fake_provider, vision_fallback=True).load_to_markdown(sample_pdf)
    assert "expected sentence from sample.pdf" in md
    fake_provider.describe_image.assert_not_called()


# ----- DocxLoader -----


@pytest.mark.unit
def test_docx_loader_preserves_headings(sample_docx: Path) -> None:
    """H1 and H2 from the source document survive as ``#`` and ``##`` in the markdown."""
    md = DocxLoader().load_to_markdown(sample_docx)
    assert "# Top-Level Title" in md
    assert "## Second-Level Section" in md


@pytest.mark.unit
def test_docx_loader_preserves_paragraph_text(sample_docx: Path) -> None:
    md = DocxLoader().load_to_markdown(sample_docx)
    assert "First paragraph under the H1." in md
    assert "Body paragraph under the H2." in md


@pytest.mark.unit
def test_docx_loader_converts_tables_to_markdown(sample_docx: Path) -> None:
    """A DOCX table becomes a pipe-delimited markdown table."""
    md = DocxLoader().load_to_markdown(sample_docx)
    assert "| header-a | header-b |" in md
    assert "| value-a | value-b |" in md


@pytest.mark.unit
def test_docx_loader_routed_via_registry(sample_docx: Path) -> None:
    assert isinstance(get_loader_for(sample_docx), DocxLoader)


@pytest.mark.unit
def test_docx_loader_can_handle_docx_only(tmp_path: Path) -> None:
    loader = DocxLoader()
    src = tmp_path / "x.docx"
    src.touch()
    assert loader.can_handle(src) is True
    assert loader.can_handle(Path("x.doc")) is False  # legacy .doc not supported (different format)
    assert loader.can_handle(Path("x.txt")) is False
